"""Excel template + import for products of one category.

Row 1 holds column titles. Feature columns end with "[key]", so labels can be renamed safely.
Rows with a SKU that already exists update that product; other rows create new products.
Nothing is saved unless every row is valid.
"""

import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.catalog.brands import brand_key, brand_lookup, resolve_brand
from app.modules.catalog.models import AttributeDefinition, Product, ProductUnit
from app.modules.catalog.schemas import UnitIn
from app.modules.catalog.service import refresh_search_words, validate_attributes
from app.modules.inventory.models import Stock
from app.modules.inventory.service import change_stock, stock_map

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
PACKS = 3

BASE_COLUMNS = {
    "name": "Name *",
    "brand": "Brand",
    "sku": "SKU",
    "mrp": "MRP",
    "description": "Description",
    "search_aliases": "Search words (Hindi / Hinglish)",
    "status": "Status (active / hidden)",
    "size_group": "Size group (same text = sizes of one product)",
    "variant_label": "Size label (e.g. 5 kg)",
}
# The two size columns are optional: old sheets without them still import and leave existing groups alone.
PACK_FIELDS = {
    "name": "name *",
    "pieces_per_unit": "pieces *",
    "price": "price *",
    "min_qty": "min qty",
    "step_qty": "step",
    "max_qty": "max qty",
}
TAIL_COLUMNS = {"default_pack": "Default pack (1-3)", "stock_pieces": "Stock (pieces)"}

_KEY_IN_HEADER = re.compile(r"\[([a-z][a-z0-9_]*)\]\s*$")


def _feature_header(d: AttributeDefinition) -> str:
    unit = f" ({d.unit})" if d.unit else ""
    star = " *" if d.required else ""
    return f"{d.label}{unit}{star} [{d.key}]"


def _pack_header(n: int, field_name: str) -> str:
    label = PACK_FIELDS[field_name]
    if n > 1:
        label = label.replace(" *", "")
    return f"Pack {n} {label}"


def build_template(category_path: list[str], defs: list[AttributeDefinition]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = list(BASE_COLUMNS.values()) + [_feature_header(d) for d in defs]
    headers += [_pack_header(n, f) for n in range(1, PACKS + 1) for f in PACK_FIELDS]
    headers += list(TAIL_COLUMNS.values())
    ws.append(headers)

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(30, len(title) + 2))
    ws.freeze_panes = "B2"

    # Dropdowns for select / yes-no features and status
    first_feature_col = len(BASE_COLUMNS) + 1
    for i, d in enumerate(defs):
        options = d.options if d.type == "select" else ["yes", "no"] if d.type == "boolean" else None
        if options:
            letter = get_column_letter(first_feature_col + i)
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
            ws.add_data_validation(dv)
    status_letter = get_column_letter(list(BASE_COLUMNS).index("status") + 1)
    dv = DataValidation(type="list", formula1='"active,hidden"', allow_blank=True)
    dv.add(f"{status_letter}2:{status_letter}{MAX_ROWS + 1}")
    ws.add_data_validation(dv)

    help_ws = wb.create_sheet("How to fill")
    lines = [
        f"Category: {' > '.join(category_path)}",
        "",
        "One row per product. Columns with * are required.",
        "Pack sizes: 'pieces' = how many pieces in that pack (Piece = 1, Box = 12, Carton = 24). Price is for the whole pack.",
        "Pack 1 is required; Pack 2 and 3 are optional. 'Default pack' is the one shops see first (1, 2 or 3).",
        "Stock is counted in pieces. On a new product it becomes the opening stock; on an existing product it sets the stock to this number.",
        "Search words: other names shops may type, e.g. cheeni, chini, शक्कर.",
        "SKU: your own product code. If a row's SKU already exists, that product is updated instead of added again.",
        "Do not change the column titles in row 1.",
        "",
        "Example row: Mithas Sugar 1 kg | Mithas | MIT-SUG-1 | 50 | | cheeni, chini | active | 1 | Medium | Packet | 1 | 45 | 2 | | | Bag | 30 | 1290 | ... | 2 | 300",
    ]
    for line in lines:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 120

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class RowData:
    row: int
    fields: dict[str, Any]
    attributes: dict[str, Any]
    units: list[UnitIn]
    default_pack: int
    stock_pieces: int | None
    existing: Product | None = None


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list[dict] = field(default_factory=list)
    new_brands: list[str] = field(default_factory=list)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _int(value: Any, label: str, errors: list[str]) -> int | None:
    text = _text(value)
    if text is None:
        return None
    try:
        number = float(text)
        if not number.is_integer():
            raise ValueError
        return int(number)
    except ValueError:
        errors.append(f"{label} must be a whole number")
        return None


def _money(value: Any, label: str, errors: list[str]) -> Decimal | None:
    text = _text(value)
    if text is None:
        return None
    try:
        amount = Decimal(text.replace(",", "").replace("₹", "").strip())
        if amount < 0:
            raise InvalidOperation
        return amount.quantize(Decimal("0.01"))
    except InvalidOperation:
        errors.append(f"{label} must be an amount like 45 or 45.50")
        return None


def _map_columns(header: list[Any], defs: list[AttributeDefinition]) -> tuple[dict[str, int], dict[str, int]]:
    """Returns (fixed column -> index, feature key -> index). Raises 422 if the sheet doesn't match."""
    titles = [(_text(h) or "") for h in header]
    expected = {**{k: v for k, v in BASE_COLUMNS.items()}, **TAIL_COLUMNS}
    for n in range(1, PACKS + 1):
        for f in PACK_FIELDS:
            expected[f"pack{n}_{f}"] = _pack_header(n, f)
    fixed: dict[str, int] = {}
    for key, title in expected.items():
        if title in titles:
            fixed[key] = titles.index(title)
    missing = [title for key, title in expected.items() if key not in fixed and key in ("name", "pack1_name", "pack1_pieces_per_unit", "pack1_price")]
    if missing:
        raise HTTPException(422, f"This doesn't look like our template. Missing columns: {', '.join(missing)}")
    features: dict[str, int] = {}
    known = {d.key for d in defs}
    for i, title in enumerate(titles):
        match = _KEY_IN_HEADER.search(title)
        if match:
            if match.group(1) not in known:
                raise HTTPException(422, f"Column '{title}' is not a feature of this category. Download a fresh template.")
            features[match.group(1)] = i
    return fixed, features


def _parse_row(row_no: int, cells: list[Any], fixed: dict[str, int], features: dict[str, int], defs) -> tuple[RowData | None, list[str]]:
    def cell(key):
        i = fixed.get(key)
        return cells[i] if i is not None and i < len(cells) else None

    errors: list[str] = []
    name = _text(cell("name"))
    if not name:
        errors.append("Name is required")
    status = (_text(cell("status")) or "active").lower()
    if status not in ("active", "hidden"):
        errors.append("Status must be active or hidden")
    fields = {
        "name": name,
        "brand": _text(cell("brand")),
        "sku": _text(cell("sku")),
        "mrp": _money(cell("mrp"), "MRP", errors),
        "description": _text(cell("description")),
        "search_aliases": _text(cell("search_aliases")),
        "status": status,
    }
    # Blank size cells leave the product as it is (groups made in the app stay).
    group = _text(cell("size_group"))
    if group:
        # Same text -> same group id, so importing the sheet again keeps the groups (queries are per tenant).
        fields["variant_group_id"] = uuid.uuid5(uuid.NAMESPACE_URL, f"orderdesk-size-group:{group.lower()}")
    label = _text(cell("variant_label"))
    if label:
        fields["variant_label"] = label[:40]

    values = {k: cells[i] for k, i in features.items() if i < len(cells) and _text(cells[i]) is not None}
    attributes: dict[str, Any] = {}
    try:
        attributes = validate_attributes(defs, values)
    except HTTPException as exc:
        errors.append(exc.detail)

    units: list[UnitIn] = []
    for n in range(1, PACKS + 1):
        raw = {f: cell(f"pack{n}_{f}") for f in PACK_FIELDS}
        if all(_text(v) is None for v in raw.values()):
            continue
        pack_errors: list[str] = []
        data = {
            "name": _text(raw["name"]),
            "pieces_per_unit": _int(raw["pieces_per_unit"], f"Pack {n} pieces", pack_errors),
            "price": _money(raw["price"], f"Pack {n} price", pack_errors),
            "min_qty": _int(raw["min_qty"], f"Pack {n} min qty", pack_errors) or 1,
            "step_qty": _int(raw["step_qty"], f"Pack {n} step", pack_errors) or 1,
            "max_qty": _int(raw["max_qty"], f"Pack {n} max qty", pack_errors),
        }
        if pack_errors:
            errors.extend(pack_errors)
            continue
        missing = [label for key, label in (("name", "name"), ("pieces_per_unit", "pieces"), ("price", "price")) if data[key] is None]
        if missing:
            errors.append(f"Pack {n} needs {', '.join(missing)}")
            continue
        try:
            units.append(UnitIn(**data))
        except ValidationError as exc:
            errors.append(f"Pack {n}: " + "; ".join(e["msg"] for e in exc.errors()))
    if not units and not any("Pack" in e for e in errors):
        errors.append("At least Pack 1 is required")
    if len({u.name.lower() for u in units}) != len(units):
        errors.append("Pack names must be different")

    default_pack = _int(cell("default_pack"), "Default pack", errors) or 1
    if units and not 1 <= default_pack <= len(units):
        errors.append(f"Default pack must be between 1 and {len(units)}")
    stock = _int(cell("stock_pieces"), "Stock", errors)
    if stock is not None and stock < 0:
        errors.append("Stock cannot be negative")

    if errors:
        return None, errors
    return RowData(row_no, fields, attributes, units, default_pack, stock), []


async def import_products(
    db: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    category_id: uuid.UUID,
    defs: list[AttributeDefinition],
    data: bytes,
    user_id: uuid.UUID,
    dry_run: bool,
) -> ImportResult:
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(413, "File is too large (max 5 MB)")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "Could not read the file. Upload the .xlsx template you downloaded.") from None
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(422, "The sheet is empty")
    fixed, features = _map_columns(list(rows[0]), defs)
    body = [(i + 2, list(r)) for i, r in enumerate(rows[1:]) if any(_text(c) is not None for c in r)]
    if not body:
        raise HTTPException(422, "No product rows found")
    if len(body) > MAX_ROWS:
        raise HTTPException(422, f"Too many rows (max {MAX_ROWS})")

    result = ImportResult()
    parsed: list[RowData] = []
    for row_no, cells in body:
        data_row, errors = _parse_row(row_no, cells, fixed, features, defs)
        if errors:
            result.errors.append({"row": row_no, "message": "; ".join(errors)})
        else:
            parsed.append(data_row)

    # SKU: unique within the file; existing SKUs mean "update"
    seen: dict[str, int] = {}
    for r in parsed:
        sku = r.fields["sku"]
        if sku:
            key = sku.lower()
            if key in seen:
                result.errors.append({"row": r.row, "message": f"SKU {sku} is also used in row {seen[key]}"})
            seen[key] = r.row
    skus = [r.fields["sku"] for r in parsed if r.fields["sku"]]
    existing = {}
    if skus:
        found = await db.scalars(select(Product).where(Product.tenant_id == tenant_id, Product.sku.in_(skus)))
        existing = {p.sku.lower(): p for p in found}
    no_sku_names = [r.fields["name"].lower() for r in parsed if not r.fields["sku"]]
    clashing_names = set()
    if no_sku_names:
        clashing_names = set(
            await db.scalars(
                select(func.lower(Product.name)).where(
                    Product.tenant_id == tenant_id, func.lower(Product.name).in_(no_sku_names)
                )
            )
        )
    for r in parsed:
        if r.fields["sku"]:
            r.existing = existing.get(r.fields["sku"].lower())
        elif r.fields["name"].lower() in clashing_names:
            result.errors.append(
                {"row": r.row, "message": f"'{r.fields['name']}' already exists. Add its SKU to update it, or rename it."}
            )

    result.errors.sort(key=lambda e: e["row"])
    if result.errors:
        return result
    result.created = sum(1 for r in parsed if r.existing is None)
    result.updated = len(parsed) - result.created
    lookup = await brand_lookup(db, tenant_id)
    result.new_brands = sorted(
        {" ".join(r.fields["brand"].split()) for r in parsed if r.fields.get("brand") and brand_key(r.fields["brand"]) not in lookup},
        key=str.lower,
    )
    if dry_run:
        return result
    for r in parsed:  # same brand however it was typed
        brand = await resolve_brand(db, tenant_id, r.fields.get("brand"), lookup)
        r.fields["brand"] = brand.name if brand else None
        r.fields["brand_id"] = brand.id if brand else None

    for r in parsed:
        if r.existing is None:
            await _create(db, tenant_id, category_id, r, user_id)
        else:
            await _update(db, tenant_id, category_id, r, user_id)
    await refresh_search_words(db, tenant_id)
    await db.commit()
    return result


def _apply_default(units: list[ProductUnit], chosen: ProductUnit) -> None:
    for u in units:
        u.is_default = u is chosen


async def _create(db, tenant_id, category_id, r: RowData, user_id) -> None:
    product = Product(tenant_id=tenant_id, category_id=category_id, attributes=r.attributes, **r.fields)
    product.units = [
        ProductUnit(tenant_id=tenant_id, sort_order=i, **u.model_dump(exclude={"is_default", "sort_order"}))
        for i, u in enumerate(r.units)
    ]
    _apply_default(product.units, product.units[r.default_pack - 1])
    db.add(product)
    await db.flush()
    db.add(Stock(product_id=product.id, tenant_id=tenant_id, qty_pieces=0))
    await db.flush()
    if r.stock_pieces:
        await change_stock(
            db, tenant_id=tenant_id, product_id=product.id, change_pieces=r.stock_pieces,
            reason="opening", user_id=user_id, note="Excel import",
        )


async def _update(db, tenant_id, category_id, r: RowData, user_id) -> None:
    product = r.existing
    for key, value in r.fields.items():
        setattr(product, key, value)
    product.category_id = category_id
    product.attributes = r.attributes
    by_name = {u.name.lower(): u for u in product.units}
    chosen = None
    for i, u in enumerate(r.units):
        unit = by_name.get(u.name.lower())
        values = u.model_dump(exclude={"is_default", "sort_order"})
        if unit is None:
            unit = ProductUnit(tenant_id=tenant_id, product_id=product.id, sort_order=i, **values)
            product.units.append(unit)
        else:
            for key, value in values.items():
                setattr(unit, key, value)
        if i == r.default_pack - 1:
            chosen = unit
    _apply_default(product.units, chosen)
    await db.flush()
    if r.stock_pieces is not None:
        current = (await stock_map(db, [product.id])).get(product.id, 0)
        if r.stock_pieces != current:
            await change_stock(
                db, tenant_id=tenant_id, product_id=product.id, change_pieces=r.stock_pieces - current,
                reason="adjust", user_id=user_id, note="Excel import",
            )
