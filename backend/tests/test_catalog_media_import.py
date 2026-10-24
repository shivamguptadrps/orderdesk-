"""M2: product photos and Excel import."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from app.core.media import media_root
from tests.conftest import sugar

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _png(width=1600, height=1200, color=(200, 50, 50, 255)) -> bytes:
    buf = BytesIO()
    Image.new("RGBA", (width, height), color).save(buf, "PNG")
    return buf.getvalue()


async def _product(client, c):
    res = await client.post(f"/t/{c['a'].id}/products", json=sugar(c["ids"]["Sugar"]), headers=c["owner"])
    assert res.status_code == 201, res.text
    return res.json()


def _file_for(url: str) -> Path:
    return media_root() / url.removeprefix("/media/")


# ---------- photos ----------


async def test_upload_resizes_and_serves(client, cat):
    p = await _product(client, cat)
    url = f"/t/{cat['a'].id}/products/{p['id']}/images"
    res = await client.post(url, files={"file": ("photo.png", _png(), "image/png")}, headers=cat["owner"])
    assert res.status_code == 201, res.text
    body = res.json()
    big, small = body["images"][0], body["thumbnails"][0]
    assert big.endswith("_800.jpg") and small.endswith("_200.jpg")
    assert Image.open(_file_for(big)).size == (800, 600)
    assert Image.open(_file_for(small)).size == (200, 150)

    served = await client.get(big)
    assert served.status_code == 200 and served.headers["content-type"] == "image/jpeg"

    card = (await client.get(f"/t/{cat['a'].id}/catalog/products", headers=cat["shop"])).json()["items"][0]
    assert card["image"] == small


async def test_reject_non_image(client, cat):
    p = await _product(client, cat)
    url = f"/t/{cat['a'].id}/products/{p['id']}/images"
    res = await client.post(url, files={"file": ("x.jpg", b"not an image at all", "image/jpeg")}, headers=cat["owner"])
    assert res.status_code == 422


async def test_delete_photo_removes_files(client, cat):
    p = await _product(client, cat)
    base = f"/t/{cat['a'].id}/products/{p['id']}/images"
    for color in [(255, 0, 0, 255), (0, 255, 0, 255)]:
        await client.post(base, files={"file": ("p.png", _png(color=color), "image/png")}, headers=cat["owner"])
    product = (await client.get(f"/t/{cat['a'].id}/products/{p['id']}", headers=cat["owner"])).json()
    first, second = product["images"]

    res = await client.delete(f"{base}/0", headers=cat["owner"])
    assert res.json()["images"] == [second]
    assert not _file_for(first).exists() and _file_for(second).exists()
    assert (await client.delete(f"{base}/5", headers=cat["owner"])).status_code == 404


async def test_photos_tenant_isolation(client, cat):
    p = await _product(client, cat)
    b = cat["b"]
    res = await client.post(
        f"/t/{b.id}/products/{p['id']}/images", files={"file": ("p.png", _png(), "image/png")}, headers=cat["owner_b"]
    )
    assert res.status_code == 404
    assert (await client.delete(f"/t/{b.id}/products/{p['id']}/images/0", headers=cat["owner_b"])).status_code == 404
    res = await client.post(
        f"/t/{cat['a'].id}/products/{p['id']}/images", files={"file": ("p.png", _png(), "image/png")}, headers=cat["shop"]
    )
    assert res.status_code == 403


# ---------- Excel ----------


async def _template(client, c, category="Sugar") -> Workbook:
    res = await client.get(
        f"/t/{c['a'].id}/product-import/template", params={"category_id": c["ids"][category]}, headers=c["owner"]
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == XLSX
    return load_workbook(BytesIO(res.content))


def _fill(wb: Workbook, rows: list[dict]) -> bytes:
    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROW = {
    "Name *": "Mithas Sugar 1 kg",
    "Brand": "Mithas",
    "SKU": "MIT-1",
    "MRP": 50,
    "Search words (Hindi / Hinglish)": "cheeni, chini",
    "Pack weight (kg) * [weight_kg]": 1,
    "Grain [grain]": "medium",
    "Pack 1 name *": "Packet",
    "Pack 1 pieces *": 1,
    "Pack 1 price *": 45,
    "Pack 1 min qty": 2,
    "Pack 2 name": "Bag",
    "Pack 2 pieces": 30,
    "Pack 2 price": "1,290",
    "Default pack (1-3)": 2,
    "Stock (pieces)": 300,
}


async def _import(client, c, content: bytes, dry_run: bool, category="Sugar"):
    res = await client.post(
        f"/t/{c['a'].id}/product-import",
        params={"category_id": c["ids"][category], "dry_run": str(dry_run).lower()},
        files={"file": ("p.xlsx", content, XLSX)},
        headers=c["owner"],
    )
    return res


async def test_template_has_feature_columns_and_dropdowns(client, cat):
    wb = await _template(client, cat)
    headers = [c.value for c in wb["Products"][1]]
    assert "Pack weight (kg) * [weight_kg]" in headers and "Grain [grain]" in headers
    validations = [dv.formula1 for dv in wb["Products"].data_validations.dataValidation]
    assert '"Fine,Medium,Bold"' in validations
    assert "How to fill" in wb.sheetnames


async def test_import_dry_run_then_save(client, cat):
    rows = [ROW, ROW | {"Name *": "Mithas Sugar 5 kg", "SKU": "MIT-5", "Pack weight (kg) * [weight_kg]": 5, "Default pack (1-3)": None}]
    content = _fill(await _template(client, cat), rows)

    res = await _import(client, cat, content, dry_run=True)
    assert res.json() == {"dry_run": True, "created": 2, "updated": 0, "errors": [], "saved": False}
    listed = (await client.get(f"/t/{cat['a'].id}/products", headers=cat["owner"])).json()
    assert listed["total"] == 0  # nothing saved on a dry run

    res = await _import(client, cat, content, dry_run=False)
    assert res.json()["saved"] is True
    listed = (await client.get(f"/t/{cat['a'].id}/products", params={"q": "mithas"}, headers=cat["owner"])).json()
    one = next(p for p in listed["items"] if p["sku"] == "MIT-1")
    assert one["attributes"] == {"weight_kg": 1, "grain": "Medium"}
    assert one["stock_pieces"] == 300
    assert {u["name"]: (u["price"], u["is_default"]) for u in one["units"]} == {"Packet": (45.0, False), "Bag": (1290.0, True)}
    five = next(p for p in listed["items"] if p["sku"] == "MIT-5")
    assert next(u for u in five["units"] if u["is_default"])["name"] == "Packet"  # default pack 1 when blank

    # shops can find it by Hinglish
    found = (await client.get(f"/t/{cat['a'].id}/catalog/products", params={"q": "chini"}, headers=cat["shop"])).json()
    assert found["total"] == 2


async def test_import_updates_by_sku(client, cat):
    await _import(client, cat, _fill(await _template(client, cat), [ROW]), dry_run=False)
    changed = ROW | {"Pack 2 price": 1250, "Pack 3 name": "Carton", "Pack 3 pieces": 50, "Pack 3 price": 2000, "Stock (pieces)": 250}
    res = await _import(client, cat, _fill(await _template(client, cat), [changed]), dry_run=False)
    assert res.json()["created"] == 0 and res.json()["updated"] == 1

    product = (await client.get(f"/t/{cat['a'].id}/products", params={"q": "mithas"}, headers=cat["owner"])).json()["items"][0]
    assert {u["name"]: u["price"] for u in product["units"]} == {"Packet": 45.0, "Bag": 1250.0, "Carton": 2000.0}
    assert product["stock_pieces"] == 250
    ledger = (await client.get(f"/t/{cat['a'].id}/stock/ledger", params={"product_id": product["id"]}, headers=cat["owner"])).json()
    assert [(e["reason"], e["change_pieces"]) for e in ledger] == [("adjust", -50), ("opening", 300)]


async def test_import_reports_every_bad_row_and_saves_nothing(client, cat):
    rows = [
        ROW,
        ROW | {"SKU": "MIT-1"},  # duplicate SKU in file
        ROW | {"SKU": "B-1", "Name *": None},
        ROW | {"SKU": "B-2", "Grain [grain]": "Huge", "Pack 1 price *": "abc"},
        ROW | {"SKU": "B-3", "Pack 1 name *": None, "Pack 1 pieces *": None, "Pack 1 price *": None, "Pack 1 min qty": None,
               "Pack 2 name": None, "Pack 2 pieces": None, "Pack 2 price": None},
        ROW | {"SKU": "B-4", "Default pack (1-3)": 3},
    ]
    res = await _import(client, cat, _fill(await _template(client, cat), rows), dry_run=False)
    body = res.json()
    assert body["saved"] is False
    by_row = {e["row"]: e["message"] for e in body["errors"]}
    assert set(by_row) == {3, 4, 5, 6, 7}
    assert "also used in row 2" in by_row[3]
    assert "Name is required" in by_row[4]
    assert "Fine, Medium, Bold" in by_row[5] and "Pack 1 price" in by_row[5]
    assert "Pack 1 is required" in by_row[6]
    assert "Default pack must be between 1 and 2" in by_row[7]
    assert (await client.get(f"/t/{cat['a'].id}/products", headers=cat["owner"])).json()["total"] == 0


async def test_import_rejects_wrong_files(client, cat):
    res = await _import(client, cat, b"hello", dry_run=True)
    assert res.status_code == 422

    wb = Workbook()
    wb.active.append(["Something", "Else"])
    buf = BytesIO()
    wb.save(buf)
    res = await _import(client, cat, buf.getvalue(), dry_run=True)
    assert res.status_code == 422 and "Missing columns" in res.json()["detail"]

    # Sugar template uploaded into Oil: grain is not an Oil feature
    content = _fill(await _template(client, cat), [ROW])
    res = await _import(client, cat, content, dry_run=True, category="Oil")
    assert res.status_code == 422


async def test_import_name_clash_without_sku(client, cat):
    await _import(client, cat, _fill(await _template(client, cat), [ROW | {"SKU": None}]), dry_run=False)
    res = await _import(client, cat, _fill(await _template(client, cat), [ROW | {"SKU": None}]), dry_run=True)
    assert "already exists" in res.json()["errors"][0]["message"]


async def test_import_tenant_isolation(client, cat):
    content = _fill(await _template(client, cat), [ROW])
    b = cat["b"]
    res = await client.get(f"/t/{b.id}/product-import/template", params={"category_id": cat["ids"]["Sugar"]}, headers=cat["owner_b"])
    assert res.status_code == 404
    res = await client.post(
        f"/t/{b.id}/product-import", params={"category_id": cat["ids"]["Sugar"], "dry_run": "false"},
        files={"file": ("p.xlsx", content, XLSX)}, headers=cat["owner_b"],
    )
    assert res.status_code == 404

