"""Load a prepared catalog (master.json from a data folder) into one distributor through the public API.

Same steps for the demo distributor and the real one, so the work is done once:
  1. brands: add missing ones, add other spellings to existing ones
  2. categories: create the paths the catalog uses if missing (matched by the English part of the name)
  3. old products: rename their SKU to the Tally name so the import updates them instead of copying
  4. Excel: one sheet per category, built from THIS distributor's own template; dry run all, then save
  5. photos: <data>/photos/<n>_<k>.jpg -> uploaded to products that have fewer than 2 photos

usage:
  python tools/catalog_load.py <data folder> <api base> <tenant id> [--dry-run] [--skip-photos]
  ORDERDESK_TOKEN=<owner or staff access token> must be set.
"""

import json
import os
import re
import sys
import urllib.error
import urllib.request
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


class Api:
    def __init__(self, base: str, tenant: str, token: str):
        self.base, self.t, self.token = base.rstrip("/"), f"/t/{tenant}", token

    def call(self, path: str, body=None, method: str | None = None, raw: bytes | None = None, ctype: str | None = None):
        data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
        req = urllib.request.Request(
            self.base + path, data=data, method=method or ("POST" if data is not None else "GET"),
            headers={"Authorization": f"Bearer {self.token}", "Content-Type": ctype or "application/json"},
        )
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                out = r.read()
                return json.loads(out) if out and r.headers.get("content-type", "").startswith("application/json") else out
        except urllib.error.HTTPError as e:
            raise SystemExit(f"{method or 'GET'} {path} -> {e.code}: {e.read().decode()[:400]}") from None

    def upload(self, path: str, filename: str, data: bytes, ctype: str):
        bnd = uuid.uuid4().hex
        body = (f'--{bnd}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
                f"Content-Type: {ctype}\r\n\r\n").encode() + data + f"\r\n--{bnd}--\r\n".encode()
        return self.call(path, raw=body, method="POST", ctype=f"multipart/form-data; boundary={bnd}")


def en(name: str) -> str:
    return name.split(" / ")[0].strip().lower()


def ensure_brands(api: Api, brands: dict) -> None:
    have = {b["name"].lower(): b for b in api.call(f"{api.t}/brands")}
    for name, aliases in brands.items():
        b = have.get(name.lower())
        if b is None:
            api.call(f"{api.t}/brands", {"name": name, "aliases": aliases})
            print(f"  + brand {name}")
            continue
        old = [a.strip() for a in (b["aliases"] or "").split(",") if a.strip()]
        new = [a.strip() for a in aliases.split(",") if a.strip() and a.strip().lower() not in {o.lower() for o in old}]
        if new:
            api.call(f"{api.t}/brands/{b['id']}", {"aliases": ", ".join(old + new)}, "PUT")
            print(f"  ~ brand {name}: + {', '.join(new)}")


def category_ids(api: Api, wanted_new: list[dict], used_paths: list[list[str]]) -> dict[tuple, str]:
    """Map each English path (Packet food, Baking) to a category id, creating what's missing."""
    full = {tuple(c["path"]): c for c in wanted_new}
    cats = api.call(f"{api.t}/categories")

    def find(parent, name):
        return next((c for c in cats if c["parent_id"] == parent and en(c["name"]) == name.lower()), None)

    out = {}
    for path in sorted({tuple(p) for p in used_paths}, key=len):
        parent = None
        for depth in range(1, len(path) + 1):
            sub = path[:depth]
            c = find(parent, sub[-1])
            if c is None:
                spec = full.get(sub, {"name": sub[-1], "aliases": None})
                c = api.call(f"{api.t}/categories", {"name": spec["name"], "aliases": spec["aliases"], "parent_id": parent})
                cats.append(c)
                print(f"  + category {' > '.join(sub)}")
            parent = c["id"]
        out[path] = parent
    return out


def rename_old_skus(api: Api, mapping: dict) -> None:
    for old, tally in mapping.items():
        found = api.call(f"{api.t}/products?q=&limit=200")["items"]
        p = next((x for x in found if (x["sku"] or "") == old), None)
        if p is None:
            continue
        if any((x["sku"] or "") == tally[:64] for x in found):
            print(f"  ! SKU {tally[:64]} already exists; left product {old} as it is")
            continue
        api.call(f"{api.t}/products/{p['id']}", {"sku": tally[:64]}, "PUT")
        print(f"  ~ product {p['name']}: SKU {old} -> {tally[:64]}")


def build_sheet(api: Api, cat_id: str, products: list[dict]) -> bytes:
    raw = api.call(f"{api.t}/product-import/template?category_id={cat_id}")
    wb = load_workbook(BytesIO(raw))
    ws = wb["Products"]
    headers = [c.value or "" for c in ws[1]]
    for p in products:
        row = []
        for h in headers:
            v = None
            if h.startswith("Name"):
                v = p["name"]
            elif h == "Brand":
                v = p["brand"]
            elif h == "SKU":
                v = p["sku"]
            elif h == "MRP":
                v = p["mrp"]
            elif h == "Description":
                v = p["description"]
            elif h.startswith("Search words"):
                v = p["search_words"]
            elif h.startswith("Status"):
                v = p["status"]
            elif h.startswith("Default pack"):
                v = 1
            elif h.startswith("Stock"):
                v = p["stock"]
            elif m := re.search(r"\[([a-z0-9_]+)\]", h):
                v = p["features"].get(m.group(1))
            elif m := re.match(r"Pack 1 (name|pieces|price|min qty|step|max qty)", h):
                v = {"name": p["pack"]["name"], "pieces": p["pack"]["pieces"], "price": p["price"], "min qty": 1, "step": 1, "max qty": None}[m.group(1)]
            row.append(v)
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_sheet(api: Api, cat_id: str, content: bytes, dry_run: bool) -> dict:
    return api.upload(
        f"{api.t}/product-import?category_id={cat_id}&dry_run={'true' if dry_run else 'false'}",
        "catalog.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def upload_photos(api: Api, data: Path, products: list[dict]) -> None:
    photos = data / "photos"
    if not photos.exists():
        print("  (no photos folder)")
        return
    by_sku = {x["sku"]: x for x in api.call(f"{api.t}/products?limit=200")["items"] if x["sku"]}
    added = 0
    for n, p in enumerate(products):
        files = sorted(photos.glob(f"{n}_*.jpg"))
        prod = by_sku.get(p["sku"])
        if not files or prod is None:
            continue
        need = max(0, 2 - len(prod["images"]))  # top up to 2 photos; never pile on duplicates
        for f in files[:need]:
            api.upload(f"{api.t}/products/{prod['id']}/images", f.name, f.read_bytes(), "image/jpeg")
            added += 1
    print(f"  photos uploaded: {added}")


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 3 or not os.environ.get("ORDERDESK_TOKEN"):
        raise SystemExit(__doc__)
    data, base, tenant = Path(args[0]), args[1], args[2]
    dry = "--dry-run" in sys.argv
    api = Api(base, tenant, os.environ["ORDERDESK_TOKEN"])
    master = json.loads((data / "master.json").read_text(encoding="utf-8"))
    products = master["products"]
    if "--photos-only" in sys.argv:  # add photos later without re-importing (keeps live stock as it is)
        print("5. photos")
        upload_photos(api, data, products)
        return

    print("1. brands")
    ensure_brands(api, master["brands"]) if not dry else print("  (dry run: skipped)")
    print("2. categories")
    ids = category_ids(api, master["categories"], [p["cat"] for p in products]) if not dry else {}
    print("3. old products")
    rename_old_skus(api, master.get("existing_sku", {})) if not dry else print("  (dry run: skipped)")
    if dry:
        return
    print("4. Excel import")
    groups: dict[str, list] = {}
    for p in products:
        groups.setdefault(ids[tuple(p["cat"])], []).append(p)
    sheets = {cid: build_sheet(api, cid, ps) for cid, ps in groups.items()}
    problems = []
    for cid, content in sheets.items():
        r = import_sheet(api, cid, content, dry_run=True)
        if r["errors"]:
            problems.append((cid, r["errors"]))
    if problems:
        raise SystemExit(f"Dry run found problems, nothing saved: {json.dumps(problems, ensure_ascii=False)[:2000]}")
    created = updated = 0
    for cid, content in sheets.items():
        r = import_sheet(api, cid, content, dry_run=False)
        created, updated = created + r["created"], updated + r["updated"]
        (data / "sheets_out").mkdir(exist_ok=True)
        (data / "sheets_out" / f"{cid}.xlsx").write_bytes(content)
    print(f"  created {created}, updated {updated} (sheets saved in {data / 'sheets_out'})")
    if "--skip-photos" not in sys.argv:
        print("5. photos")
        upload_photos(api, data, products)


if __name__ == "__main__":
    main()
